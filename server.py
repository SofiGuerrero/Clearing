"""
SAP Clearing Dashboard — servidor local Flask
Requiere: pip install flask
"""

import os
import sys
import re
import io
import time
import queue
import shutil
import tempfile
import threading
import subprocess
import openpyxl
from datetime import datetime
from flask import Flask, Response, render_template_string, jsonify, request, send_file
import pandas as pd

app = Flask(__name__)

SCRIPT_DIR = os.path.join(
    os.path.expanduser("~"),
    "OneDrive - SAP SE", "Varios SAP - Daily", "JE-Sofi", "Clearing"
)

SCRIPTS = {
    "us":     "Clearing US Complete v2.py",
    "canada": "Clearing Canada Complete v2.py",
    "concur": "Clearing Concur Complete v2.py",
}

OUTPUT_FILES = {
    "us":     "Clearing US Pivot.xlsx",
    "canada": "Clearing Canada Pivot.xlsx",
    "concur": "Clearing Concur Pivot.xlsx",
}

INPUT_FILES = {
    "us":     "Clearing US Input.xlsx",
    "canada": "Clearing Canada Input.xlsx",
    "concur": "Clearing Concur Input.xlsx",
}

_clearing_results = {}    # key → [{account, cc, clearing_doc, status, error}]
_clear_tasks = {}         # task_id → {"key": ..., "q": Queue}
_clear_pause_events = {}  # task_id → threading.Event (waiting for manual SAP action)


def do_clear_account(session, gl_account, company_code, doc_numbers, clearing_date, system="ISP"):
    """Execute SAP F-03 clearing. Returns (clearing_doc_number, error_message).
    system='ISP' (US/Canada) uses radio[3,0] and wnd[1] for the doc number.
    system='I4P' (Concur) uses radio[2,0] and reads the doc number from the status bar.
    NOTE: radio button indices may differ per SAP GUI template -- adjust if needed.
    """
    try:
        session.findById("wnd[0]").maximize()
        session.findById("wnd[0]/tbar[0]/okcd").text = "/nf-03"
        session.findById("wnd[0]").sendVKey(0)
        time.sleep(0.5)

        for date_path in ("wnd[0]/usr/ctxtRF05A-AUGDT", "wnd[0]/usr/ctxtBKPF-BUDAT"):
            try:
                session.findById(date_path).text = clearing_date
                break
            except Exception:
                pass

        radio_idx = 2 if system == "I4P" else 3
        radio_path = f"wnd[0]/usr/sub:SAPMF05A:0131/radRF05A-XPOS1[{radio_idx},0]"

        session.findById(radio_path).select()
        session.findById("wnd[0]/usr/ctxtRF05A-AGKON").text = str(gl_account)
        try:
            session.findById("wnd[0]/usr/ctxtBKPF-BUKRS").text = str(company_code).zfill(4)
        except Exception:
            pass
        session.findById(radio_path).setFocus()
        session.findById("wnd[0]").sendVKey(0)
        time.sleep(0.5)

        try:
            sbar = session.findById("wnd[0]/sbar").Text or ""
            if "future" in sbar.lower():
                session.findById("wnd[0]").sendVKey(0)
                time.sleep(0.5)
        except Exception:
            pass

        unique_docs = list(dict.fromkeys(str(d) for d in doc_numbers))
        chunks = [unique_docs[i:i + 19] for i in range(0, len(unique_docs), 19)]

        for chunk_idx, chunk in enumerate(chunks):
            is_last = (chunk_idx == len(chunks) - 1)
            for i, doc in enumerate(chunk):
                try:
                    session.findById(f"wnd[0]/usr/sub:SAPMF05A:0731/txtRF05A-SEL01[{i},0]").text = doc
                except Exception:
                    break
            session.findById("wnd[0]/tbar[1]/btn[16]").press()
            time.sleep(0.5)
            if not is_last:
                session.findById("wnd[0]/tbar[1]/btn[7]").press()
                time.sleep(0.3)
                session.findById("wnd[0]/tbar[1]/btn[6]").press()
                time.sleep(0.3)
                cont_radio = f"wnd[0]/usr/sub:SAPMF05A:0710/radRF05A-XPOS1[{radio_idx},0]"
                session.findById(cont_radio).select()
                session.findById(cont_radio).setFocus()
                session.findById("wnd[0]").sendVKey(0)
                time.sleep(0.5)

        try:
            sbar = session.findById("wnd[0]/sbar").Text
            if sbar and any(w in sbar.lower() for w in ("error", "not possible", "cannot")):
                return None, sbar
        except Exception:
            pass

        session.findById("wnd[0]/tbar[0]/btn[11]").press()
        time.sleep(1.0)

        if system == "I4P":
            doc_num = None
            sbar_text = ""
            try:
                sbar_text = session.findById("wnd[0]/sbar").Text or ""
                m = re.search(r"\b(\d{8,10})\b", sbar_text)
                if m:
                    doc_num = m.group(1)
            except Exception:
                pass
            try:
                session.findById("wnd[0]/sbar").doubleClick()
                time.sleep(0.5)
                if not doc_num:
                    for sp in ("wnd[0]/shellcont/shell", "wnd[0]/shellcont"):
                        try:
                            t = session.findById(sp).Text or ""
                            m2 = re.search(r"\b(\d{8,10})\b", t)
                            if m2:
                                doc_num = m2.group(1)
                                break
                        except Exception:
                            pass
            except Exception:
                pass
            try:
                session.findById("wnd[0]/shellcont").close()
                time.sleep(0.5)
            except Exception:
                pass
            if doc_num:
                return doc_num, None
            return None, f"__PAUSE_NEEDED__:{sbar_text or 'I4P: no SAP response after posting'}"

        doc_num = None
        debug_texts = []
        wnd1_found = False
        try:
            wnd1 = session.findById("wnd[1]")
            wnd1_found = True
            title = wnd1.Text or ""
            if title:
                debug_texts.append(f"title:{title}")
            try:
                usr = session.findById("wnd[1]/usr")
                for i in range(min(usr.Children.Count, 30)):
                    try:
                        child = usr.Children(i)
                        t = getattr(child, "Text", "") or ""
                        if t:
                            debug_texts.append(t)
                    except Exception:
                        pass
            except Exception:
                pass
            full_text = " ".join([title] + debug_texts)
            m = re.search(r"\b(\d{8,10})\b", full_text)
            if m:
                doc_num = m.group(1)
        except Exception:
            pass

        if wnd1_found:
            try:
                session.findById("wnd[1]/tbar[0]/btn[0]").press()
                time.sleep(1.0)
            except Exception:
                time.sleep(0.5)

        if doc_num:
            return doc_num, None

        if not wnd1_found:
            sbar = ""
            for _ in range(2):
                try:
                    sbar = session.findById("wnd[0]/sbar").Text or ""
                    if sbar:
                        break
                except Exception:
                    pass
                time.sleep(0.5)
            return None, f"__PAUSE_NEEDED__:{sbar or 'SAP rejected posting (balance != 0)'}"

        sbar = ""
        for _ in range(3):
            try:
                sbar = session.findById("wnd[0]/sbar").Text or ""
                if sbar:
                    break
            except Exception:
                pass
            time.sleep(0.8)
        m = re.search(r"\b(\d{8,10})\b", sbar)
        if m:
            return m.group(1), None

        debug = " | ".join(debug_texts) or "empty window"
        return None, f"Posted but doc number not captured. wnd[1]: {debug} | sbar: '{sbar}'"

    except Exception as exc:
        return None, str(exc)


def _retry_post(session, system="ISP"):
    """Retry F-03 post in Process Open Items screen after manual correction."""
    try:
        session.findById("wnd[0]/tbar[0]/btn[11]").press()
        time.sleep(1.0)

        if system == "I4P":
            doc_num = None
            sbar_text = ""
            try:
                sbar_text = session.findById("wnd[0]/sbar").Text or ""
                m = re.search(r"\b(\d{8,10})\b", sbar_text)
                if m:
                    doc_num = m.group(1)
            except Exception:
                pass
            try:
                session.findById("wnd[0]/sbar").doubleClick()
                time.sleep(0.5)
                if not doc_num:
                    for sp in ("wnd[0]/shellcont/shell", "wnd[0]/shellcont"):
                        try:
                            t = session.findById(sp).Text or ""
                            m2 = re.search(r"\b(\d{8,10})\b", t)
                            if m2:
                                doc_num = m2.group(1)
                                break
                        except Exception:
                            pass
            except Exception:
                pass
            try:
                session.findById("wnd[0]/shellcont").close()
                time.sleep(0.5)
            except Exception:
                pass
            if doc_num:
                return doc_num, None
            return None, f"Retry I4P: {sbar_text or 'no SAP response'}"

        doc_num = None
        debug_texts = []
        try:
            wnd1 = session.findById("wnd[1]")
            title = wnd1.Text or ""
            if title:
                debug_texts.append(f"title:{title}")
            try:
                usr = session.findById("wnd[1]/usr")
                for i in range(min(usr.Children.Count, 30)):
                    try:
                        child = usr.Children(i)
                        t = getattr(child, "Text", "") or ""
                        if t:
                            debug_texts.append(t)
                    except Exception:
                        pass
            except Exception:
                pass
            full_text = " ".join([title] + debug_texts)
            m = re.search(r"\b(\d{8,10})\b", full_text)
            if m:
                doc_num = m.group(1)
        except Exception:
            pass
        try:
            session.findById("wnd[1]/tbar[0]/btn[0]").press()
            time.sleep(1.0)
        except Exception:
            time.sleep(0.5)
        if doc_num:
            return doc_num, None
        sbar = ""
        try:
            sbar = session.findById("wnd[0]/sbar").Text or ""
        except Exception:
            pass
        m2 = re.search(r"\b(\d{8,10})\b", sbar)
        if m2:
            return m2.group(1), None
        debug = " | ".join(debug_texts)
        return None, f"Retry: {sbar or debug or 'no SAP response'}"
    except Exception as exc:
        return None, str(exc)


def _run_clearing_task(task_id, key, data):
    """Background thread: runs F-03 clearing and sends progress to the task queue."""
    import pythoncom
    pythoncom.CoInitialize()
    try:
        _do_clearing(task_id, key, data)
    finally:
        pythoncom.CoUninitialize()


def _do_clearing(task_id, key, data):
    q = _clear_tasks[task_id]["q"]
    clearing_date = data.get("clearing_date", "")
    accounts = data.get("accounts", [])

    system_label = {"us": "United States (ISP)", "canada": "Canada (ISP)", "concur": "Concur (I4P)"}.get(key, key)
    q.put(f"System: {system_label} -- {len(accounts)} account(s) -- date: {clearing_date}")

    try:
        import win32com.client
        system_prefix = "I4P" if key == "concur" else "ISP"
        sap_gui_auto = win32com.client.GetObject("SAPGUI")
        sap_app = sap_gui_auto.GetScriptingEngine
        session = None
        for i in range(sap_app.Children.Count):
            conn = sap_app.Children(i)
            if conn.Children.Count > 0:
                if conn.Children(0).PassportSystemid[:3] == system_prefix:
                    session = conn.Children(0)
                    break
        if session is None:
            q.put(f"ERROR: No {system_prefix} session found in SAP GUI")
            q.put("__DONE_CLEAR__:ERROR")
            return
    except Exception as exc:
        q.put(f"ERROR connecting SAP GUI: {exc}")
        q.put("__DONE_CLEAR__:ERROR")
        return

    results = []
    ok_count = 0

    for acc_data in accounts:
        account = str(acc_data.get("account", ""))
        cc      = str(acc_data.get("cc", ""))
        docs    = [str(d) for d in acc_data.get("documents", [])]

        q.put(f"-> Clearing {account} / {cc}  ({len(docs)} documents)...")
        clearing_doc, error = do_clear_account(session, account, cc, docs, clearing_date, system=system_prefix)

        if error and error.startswith("__PAUSE_NEEDED__:"):
            sap_msg = error[len("__PAUSE_NEEDED__:"):]
            q.put(f"WARNING SAP rejected posting for {account}/{cc}: {sap_msg}")
            q.put(f"__PAUSE__:{sap_msg}")
            pause_evt = threading.Event()
            _clear_pause_events[task_id] = pause_evt
            resumed = pause_evt.wait(timeout=300)
            _clear_pause_events.pop(task_id, None)
            if resumed:
                q.put(f"-> Retrying posting {account} / {cc}...")
                clearing_doc, error = _retry_post(session, system=system_prefix)
            else:
                error = "Timeout: manual intervention not received (5 min)"

        result = {
            "account":      account,
            "cc":           cc,
            "clearing_doc": clearing_doc or "",
            "status":       "ok" if clearing_doc else "error",
            "error":        error or "",
        }
        results.append(result)

        if clearing_doc:
            ok_count += 1
            q.put(f"OK {account} / {cc}  ->  Clearing doc: {clearing_doc}")
        else:
            q.put(f"ERROR {account} / {cc}  ->  {error}")

    _clearing_results[key] = results

    try:
        ts    = datetime.now().strftime("%Y%m%d_%H%M%S")
        fname = f"Clearing_Results_{key.upper()}_{ts}.xlsx"
        fpath = os.path.join(SCRIPT_DIR, fname)
        wb_out = openpyxl.Workbook()
        ws_out = wb_out.active
        ws_out.title = "Clearing Results"
        ws_out.append(["G/L Account", "Company Code", "Clearing Document",
                        "Status", "Error", "Clearing Date"])
        for r in results:
            ws_out.append([r["account"], r["cc"], r["clearing_doc"],
                           r["status"], r["error"], clearing_date])
        wb_out.save(fpath)
        q.put(f"Excel saved: {fname}")
    except Exception as exc:
        q.put(f"Error saving Excel: {exc}")

    q.put(f"Completed: {ok_count}/{len(accounts)} accounts cleared")
    q.put(f"__DONE_CLEAR__:{'OK' if ok_count > 0 else 'ERROR'}")


HTML = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>SAP Clearing</title>
<script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.0/dist/chart.umd.min.js"></script>
<style>
  :root {
    --shell: #1D2D3E; --blue: #0070F2; --blue-dk: #0057C2;
    --bg: #F5F6F7; --card: #FFFFFF; --text: #1D2D3E; --sub: #556B82;
    --ok: #188918; --err: #BB0000; --warn: #E9730C;
    --border: #E0E0E0; --r: 8px; --sh: 0 1px 4px rgba(0,0,0,.10);
    --font: '72','72full','SAP 72',Arial,sans-serif;
  }
  *{box-sizing:border-box;margin:0;padding:0;}
  body{font-family:var(--font);background:var(--bg);min-height:100vh;color:var(--text);}

  .shell{background:var(--shell);height:44px;display:flex;align-items:center;padding:0 16px;gap:14px;}
  .shell-brand{color:#fff;font-size:.7rem;font-weight:700;letter-spacing:1.5px;text-transform:uppercase;padding-right:14px;border-right:1px solid rgba(255,255,255,.2);}
  .shell-title{color:rgba(255,255,255,.8);font-size:.85rem;}

  .page-header{background:var(--card);border-bottom:1px solid var(--border);padding:18px 32px;display:flex;align-items:center;gap:14px;}
  .joule-badge{width:38px;height:38px;background:linear-gradient(135deg,#0070F2,#003CB4);border-radius:10px;display:flex;align-items:center;justify-content:center;color:#fff;font-size:1.1rem;flex-shrink:0;box-shadow:0 2px 8px rgba(0,112,242,.35);}
  .page-header h1{font-size:1.15rem;font-weight:700;}
  .page-header p{font-size:.78rem;color:var(--sub);margin-top:2px;}

  .content{padding:24px 32px;max-width:1200px;margin:0 auto;}
  .section-label{font-size:.7rem;font-weight:700;color:var(--sub);text-transform:uppercase;letter-spacing:.9px;margin-bottom:14px;}

  /* Execution cards */
  .exec-grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(280px,1fr));gap:14px;margin-bottom:28px;}
  .card{background:var(--card);border:1px solid var(--border);border-radius:var(--r);box-shadow:var(--sh);overflow:hidden;}
  .card-hdr{padding:14px 18px;display:flex;align-items:center;gap:12px;border-bottom:1px solid var(--border);}
  .card-icon{width:34px;height:34px;border-radius:7px;display:flex;align-items:center;justify-content:center;font-size:1rem;flex-shrink:0;}
  .icon-us{background:#EBF3FF;} .icon-ca{background:#FFF4E5;} .icon-co{background:#EDFBEE;}
  .card-hdr h2{font-size:.9rem;font-weight:700;}
  .card-hdr .sub{font-size:.72rem;color:var(--sub);margin-top:1px;}
  .card-body{padding:14px 18px;}
  .card-actions{display:flex;align-items:center;gap:10px;}
  .btn{display:inline-flex;align-items:center;gap:6px;padding:6px 16px;border:1px solid var(--blue);border-radius:4px;background:var(--blue);color:#fff;font-size:.85rem;font-weight:600;cursor:pointer;transition:background .12s;font-family:var(--font);}
  .btn:not(:disabled):hover{background:var(--blue-dk);border-color:var(--blue-dk);}
  .btn:disabled{opacity:.4;cursor:not-allowed;}
  .status{display:inline-flex;align-items:center;gap:5px;font-size:.73rem;font-weight:600;padding:3px 10px;border-radius:100px;}
  .dot{width:7px;height:7px;border-radius:50%;flex-shrink:0;}
  .status.idle{background:#EEE;color:#757575;} .status.idle .dot{background:#9E9E9E;}
  .status.running{background:#EBF3FF;color:var(--blue);} .status.running .dot{background:var(--blue);animation:blink 1s infinite;}
  .status.ok{background:#EDFBEE;color:var(--ok);} .status.ok .dot{background:var(--ok);}
  .status.error{background:#FFF0F0;color:var(--err);} .status.error .dot{background:var(--err);}
  @keyframes blink{0%,100%{opacity:1}50%{opacity:.3}}
  .log-toggle{margin-top:8px;font-size:.72rem;color:var(--blue);cursor:pointer;background:none;border:none;padding:0;font-family:var(--font);}
  .log-toggle:hover{text-decoration:underline;}
  .btn-ghost{display:inline-flex;align-items:center;gap:5px;padding:5px 12px;border:1px solid var(--border);border-radius:4px;background:#fff;color:var(--text);font-size:.82rem;font-weight:500;cursor:pointer;transition:background .12s;font-family:var(--font);}
  .btn-ghost:hover{background:var(--bg);}
  .log-wrap{margin-top:8px;border-radius:6px;overflow:hidden;display:none;}
  .log-wrap.visible{display:block;}
  .log-bar{background:#1A2535;padding:5px 12px;font-size:.66rem;color:#546E7A;font-family:Consolas,monospace;text-transform:uppercase;letter-spacing:.5px;}
  .log-body{background:#0D1520;padding:10px 12px;height:130px;overflow-y:auto;font-family:'Cascadia Code',Consolas,monospace;font-size:.73rem;color:#90A4AE;white-space:pre-wrap;word-break:break-all;}
  .line-ok{color:#66BB6A;} .line-err{color:#EF5350;} .line-inf{color:#42A5F5;}

  /* Results section */
  .results-section{display:none;}
  .results-section.visible{display:block;}
  .result-panel{background:var(--card);border:1px solid var(--border);border-radius:var(--r);box-shadow:var(--sh);margin-bottom:16px;overflow:hidden;display:none;}
  .result-panel.visible{display:block;}
  .panel-hdr{padding:14px 20px;display:flex;align-items:center;gap:12px;border-bottom:2px solid var(--border);background:#FAFBFC;}
  .panel-hdr .flag{font-size:1.2rem;}
  .panel-hdr h3{font-size:.95rem;font-weight:700;flex:1;}
  .panel-badge{background:#FFF0F0;color:var(--err);padding:3px 10px;border-radius:100px;font-size:.72rem;font-weight:700;}
  .panel-badge.all-clear{background:#EDFBEE;color:var(--ok);}

  /* Results table */
  .rtbl{width:100%;border-collapse:collapse;font-size:.82rem;}
  .rtbl th{background:#F5F6F7;color:var(--sub);font-size:.68rem;font-weight:700;text-transform:uppercase;letter-spacing:.5px;padding:8px 16px;text-align:left;border-bottom:2px solid var(--border);position:sticky;top:0;z-index:1;}
  .rtbl th.r{text-align:right;}
  .rtbl td{padding:7px 16px;border-bottom:1px solid #F2F2F2;vertical-align:middle;}
  /* Account group header row */
  .acc-hdr td{background:#EBF3FF;font-weight:700;font-size:.84rem;padding:8px 16px;border-bottom:1px solid #C8DEF7;border-top:2px solid #B8D0F0;}
  .acc-hdr .acc-num{color:var(--blue);}
  .acc-hdr .acc-bal{text-align:right;color:var(--err);}
  .acc-hdr .acc-cc{color:var(--text);}
  /* Document rows */
  .doc-row td{background:#fff;font-size:.8rem;}
  .doc-row:hover td{background:#FAFBFF;}
  .doc-row td:first-child{padding-left:32px;color:var(--sub);}
  .doc-row .amt{text-align:right;font-variant-numeric:tabular-nums;}
  .doc-row .amt.neg{color:var(--err);font-weight:600;}
  .doc-row .amt.pos{color:var(--ok);font-weight:600;}
  .doc-row .txt{color:var(--sub);max-width:200px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;}
  .rtbl tr:last-child td{border-bottom:none;}
  .tbl-scroll{max-height:500px;overflow-y:auto;}

  /* ── Analytics ─────────────────────────────────── */
  .analytics-section{margin-bottom:28px;}
  .flow-tabs{display:flex;gap:6px;flex-shrink:0;}
  .flow-tab{padding:5px 14px;border:1px solid var(--border);border-radius:100px;background:#fff;font-size:.78rem;font-weight:600;color:var(--sub);cursor:pointer;font-family:var(--font);transition:all .12s;}
  .flow-tab:hover{border-color:var(--blue);color:var(--blue);}
  .flow-tab.active{background:var(--blue);border-color:var(--blue);color:#fff;}
  .kpi-row{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin-bottom:14px;}
  @media(max-width:800px){.kpi-row{grid-template-columns:repeat(2,1fr);}}
  .kpi-tile{background:var(--card);border:1px solid var(--border);border-radius:var(--r);padding:14px 18px;box-shadow:var(--sh);}
  .kpi-tile.kpi-green{border-left:3px solid #0ca30c;}
  .kpi-tile.kpi-red{border-left:3px solid #d03b3b;}
  .kpi-tile.kpi-blue{border-left:3px solid var(--blue);}
  .kpi-label{font-size:.66rem;font-weight:700;text-transform:uppercase;letter-spacing:.8px;color:var(--sub);margin-bottom:6px;}
  .kpi-value{font-size:1.7rem;font-weight:700;color:var(--text);line-height:1;}
  .kpi-sub{font-size:.71rem;color:var(--sub);margin-top:4px;}
  .charts-row{display:grid;grid-template-columns:1fr 1fr;gap:14px;margin-bottom:14px;}
  @media(max-width:800px){.charts-row{grid-template-columns:1fr;}}
  .chart-card{background:var(--card);border:1px solid var(--border);border-radius:var(--r);box-shadow:var(--sh);padding:16px 20px;}
  .chart-card-full{margin-bottom:14px;}
  .chart-card-title{font-size:.78rem;font-weight:700;color:var(--text);margin-bottom:12px;display:flex;align-items:center;gap:8px;}
  .chart-card-title .ct-sub{font-weight:400;color:var(--sub);font-size:.72rem;}
  .chart-wrap{position:relative;}

  /* ── Clearing modal ─────────────────────────────────────────── */
  .modal-ov{position:fixed;inset:0;background:rgba(0,0,0,.45);z-index:300;align-items:center;justify-content:center;}
  .modal-box{background:#fff;border-radius:8px;width:390px;box-shadow:0 8px 32px rgba(0,0,0,.22);overflow:hidden;}
  .modal-hdr{background:var(--shell);color:#fff;padding:14px 20px;display:flex;align-items:center;justify-content:space-between;font-size:.92rem;font-weight:700;}
  .modal-x{background:none;border:none;color:#fff;cursor:pointer;font-size:1rem;opacity:.7;font-family:var(--font);line-height:1;padding:0;}
  .modal-x:hover{opacity:1;}
  .modal-body{padding:18px 20px 10px;}
  .modal-lbl{display:block;font-size:.72rem;font-weight:700;color:var(--sub);text-transform:uppercase;letter-spacing:.7px;margin-bottom:6px;}
  .modal-body input[type=date]{width:100%;padding:8px 12px;border:1px solid var(--border);border-radius:4px;font-size:.9rem;font-family:var(--font);color:var(--text);}
  .modal-log{margin:8px 20px;background:#0D1520;border-radius:6px;padding:10px 12px;height:108px;overflow-y:auto;font-family:'Cascadia Code',Consolas,monospace;font-size:.71rem;color:#90A4AE;white-space:pre-wrap;display:none;}
  .modal-log.on{display:block;}
  .modal-foot{padding:12px 20px;border-top:1px solid var(--border);display:flex;justify-content:flex-end;gap:8px;}
  .clear-btn{padding:3px 10px;border:1px solid var(--blue);border-radius:100px;background:#fff;color:var(--blue);font-size:.71rem;font-weight:600;cursor:pointer;font-family:var(--font);transition:all .1s;vertical-align:middle;}
  .clear-btn:not(:disabled):hover{background:var(--blue);color:#fff;}
  .clear-btn:disabled{opacity:.38;cursor:not-allowed;}
  .cdoc{display:inline-block;font-size:.7rem;font-weight:600;margin-left:6px;padding:2px 8px;border-radius:100px;vertical-align:middle;}
  .cdoc-ok{color:var(--ok);background:#EDFBEE;}
  .cdoc-err{color:var(--err);background:#FFF0F0;}
</style>
</head>
<body>

<div class="shell">
  <span class="shell-brand">SAP</span>
  <span class="shell-title">Clearing Automation</span>
</div>

<div class="page-header">
  <div class="joule-badge">✦</div>
  <div>
    <h1>Clearing Dashboard</h1>
    <p>FAGLL03 report automation — ISP &amp; I4P</p>
  </div>
</div>

<div class="content">
  <div class="section-label">Run Report</div>
  <div class="exec-grid">

    <div class="card">
      <div class="card-hdr">
        <div class="card-icon icon-us">🇺🇸</div>
        <div><h2>United States</h2><div class="sub">Variant /GL ACC SA · ISP</div></div>
      </div>
      <div class="card-body">
        <div class="card-actions">
          <button class="btn" onclick="run('us',this)">▶ Run</button>
          <button class="btn-ghost" onclick="loadResults('us')">↻ Load Results</button>
          <a class="btn-ghost" href="/download-open-items/us">⬇ Open Items</a>
          <span id="status-us" class="status idle"><span class="dot"></span>Idle</span>
        </div>
        <button class="log-toggle" onclick="toggleLog('us',this)" style="display:none" id="logtgl-us">Show output ▾</button>
        <div id="log-us" class="log-wrap"><div class="log-bar">output</div><div id="term-us" class="log-body"></div></div>
      </div>
    </div>

    <div class="card">
      <div class="card-hdr">
        <div class="card-icon icon-ca">🇨🇦</div>
        <div><h2>Canada</h2><div class="sub">Variant /GLCANADA · ISP</div></div>
      </div>
      <div class="card-body">
        <div class="card-actions">
          <button class="btn" onclick="run('canada',this)">▶ Run</button>
          <button class="btn-ghost" onclick="loadResults('canada')">↻ Load Results</button>
          <a class="btn-ghost" href="/download-open-items/canada">⬇ Open Items</a>
          <span id="status-canada" class="status idle"><span class="dot"></span>Idle</span>
        </div>
        <button class="log-toggle" onclick="toggleLog('canada',this)" style="display:none" id="logtgl-canada">Show output ▾</button>
        <div id="log-canada" class="log-wrap"><div class="log-bar">output</div><div id="term-canada" class="log-body"></div></div>
      </div>
    </div>

    <div class="card">
      <div class="card-hdr">
        <div class="card-icon icon-co">📋</div>
        <div><h2>Concur</h2><div class="sub">Variant PAYROLLGLCLEAR · I4P</div></div>
      </div>
      <div class="card-body">
        <div class="card-actions">
          <button class="btn" onclick="run('concur',this)">▶ Run</button>
          <button class="btn-ghost" onclick="loadResults('concur')">↻ Load Results</button>
          <a class="btn-ghost" href="/download-open-items/concur">⬇ Open Items</a>
          <span id="status-concur" class="status idle"><span class="dot"></span>Idle</span>
        </div>
        <button class="log-toggle" onclick="toggleLog('concur',this)" style="display:none" id="logtgl-concur">Show output ▾</button>
        <div id="log-concur" class="log-wrap"><div class="log-bar">output</div><div id="term-concur" class="log-body"></div></div>
      </div>
    </div>

  </div>

  <!-- Analytics section -->
  <div id="analytics-section" class="analytics-section" style="display:none;">
    <div style="display:flex;align-items:center;justify-content:space-between;margin-bottom:14px;">
      <div class="section-label" style="margin:0;">Analytics</div>
      <div class="flow-tabs">
        <button class="flow-tab" id="atab-us"     onclick="switchAnalytics('us',this)">🇺🇸 US</button>
        <button class="flow-tab" id="atab-canada" onclick="switchAnalytics('canada',this)">🇨🇦 Canada</button>
        <button class="flow-tab" id="atab-concur" onclick="switchAnalytics('concur',this)">📋 Concur</button>
      </div>
    </div>

    <!-- KPI row -->
    <div class="kpi-row">
      <div class="kpi-tile">
        <div class="kpi-label">Total Combinations</div>
        <div class="kpi-value" id="kpi-total">—</div>
        <div class="kpi-sub">Account × Company Code</div>
      </div>
      <div class="kpi-tile kpi-green">
        <div class="kpi-label">Clearable</div>
        <div class="kpi-value" id="kpi-clear">—</div>
        <div class="kpi-sub" id="kpi-clear-sub">Balance = 0</div>
      </div>
      <div class="kpi-tile kpi-red">
        <div class="kpi-label">Open Items</div>
        <div class="kpi-value" id="kpi-open">—</div>
        <div class="kpi-sub">Balance ≠ 0</div>
      </div>
      <div class="kpi-tile kpi-blue">
        <div class="kpi-label">Docs to Clear</div>
        <div class="kpi-value" id="kpi-docs">—</div>
        <div class="kpi-sub">clearable documents</div>
      </div>
    </div>

    <!-- Charts row 1 -->
    <div class="charts-row">
      <div class="chart-card">
        <div class="chart-card-title">Clear vs Open Items
          <span class="ct-sub">by Account</span>
        </div>
        <div class="chart-wrap" id="wrap-status"><canvas id="chart-status"></canvas></div>
      </div>
      <div class="chart-card">
        <div class="chart-card-title">Clearable Documents
          <span class="ct-sub">by Account</span>
        </div>
        <div class="chart-wrap" id="wrap-docs"><canvas id="chart-docs"></canvas></div>
      </div>
    </div>

    <!-- Charts row 2: period (only when Year/Month available) -->
    <div class="chart-card chart-card-full" id="period-card" style="display:none;">
      <div class="chart-card-title">Clearable Documents
        <span class="ct-sub">by Period</span>
      </div>
      <div style="position:relative;height:180px;"><canvas id="chart-period"></canvas></div>
    </div>
  </div>

  <!-- Results section -->
  <div id="results-section" class="results-section">
    <div class="section-label">Clearable Documents</div>

    <div id="panel-us" class="result-panel">
      <div class="panel-hdr">
        <span class="flag">🇺🇸</span>
        <h3>United States</h3>
        <span id="badge-us" class="panel-badge"></span>
        <button id="clearall-us" class="clear-btn" style="display:none;margin-left:auto" onclick="startClearAll('us')">&#9889; Clear All</button>
      </div>
      <div class="tbl-scroll">
        <table class="rtbl">
          <thead><tr><th>G/L Account</th><th>Company Code</th><th>Doc Number</th><th>Year/Month</th><th class="r">Amount</th><th>Text</th></tr></thead>
          <tbody id="tbody-us"></tbody>
        </table>
      </div>
    </div>

    <div id="panel-canada" class="result-panel">
      <div class="panel-hdr">
        <span class="flag">🇨🇦</span>
        <h3>Canada</h3>
        <span id="badge-canada" class="panel-badge"></span>
        <button id="clearall-canada" class="clear-btn" style="display:none;margin-left:auto" onclick="startClearAll('canada')">&#9889; Clear All</button>
      </div>
      <div class="tbl-scroll">
        <table class="rtbl">
          <thead><tr><th>G/L Account</th><th>Company Code</th><th>Doc Number</th><th>Year/Month</th><th class="r">Amount</th><th>Text</th></tr></thead>
          <tbody id="tbody-canada"></tbody>
        </table>
      </div>
    </div>

    <div id="panel-concur" class="result-panel">
      <div class="panel-hdr">
        <span class="flag">📋</span>
        <h3>Concur</h3>
        <span id="badge-concur" class="panel-badge"></span>
        <button id="clearall-concur" class="clear-btn" style="display:none;margin-left:auto" onclick="startClearAll('concur')">&#9889; Clear All</button>
      </div>
      <div class="tbl-scroll">
        <table class="rtbl">
          <thead><tr><th>G/L Account</th><th>Company Code</th><th>Doc Number</th><th>Year/Month</th><th class="r">Amount</th><th>Text</th></tr></thead>
          <tbody id="tbody-concur"></tbody>
        </table>
      </div>
    </div>

  </div>
</div>


<!-- Clearing date modal -->
<div id="clear-modal" class="modal-ov" style="display:none">
  <div class="modal-box">
    <div class="modal-hdr">
      <span>&#9889; Clear Accounts</span>
      <button class="modal-x" onclick="_closeModal()">&#10005;</button>
    </div>
    <div class="modal-body">
      <label class="modal-lbl">Clearing Date</label>
      <input type="date" id="clear-date-inp">
    </div>
    <div id="clear-modal-log" class="modal-log"></div>
    <div class="modal-foot">
      <button class="btn-ghost" id="clear-cancel-btn" onclick="_closeModal()">Cancel</button>
      <button class="btn" id="clear-resume-btn" style="display:none;background:#E9730C;border-color:#C45800;" onclick="resumeClear()">&#9654; Continue</button>
      <button class="btn" id="clear-ok-btn" onclick="confirmClear()">&#9889; Proceed</button>
    </div>
  </div>
</div>

<script>
function fmt(n){return Number(n).toLocaleString('en-US',{minimumFractionDigits:2,maximumFractionDigits:2});}

const _rdata={};
let _pend=null;
let _currentTaskId=null;

function toggleLog(key,btn){
  const w=document.getElementById('log-'+key);
  const open=w.classList.toggle('visible');
  btn.textContent=open?'Hide output ▴':'Show output ▾';
}

function run(key,btn){
  const term=document.getElementById('term-'+key);
  const status=document.getElementById('status-'+key);
  const logtgl=document.getElementById('logtgl-'+key);
  btn.disabled=true; term.innerHTML='';
  status.className='status running';
  status.innerHTML='<span class="dot"></span>Running...';
  logtgl.style.display='none';

  const es=new EventSource('/run/'+key);
  es.onmessage=function(e){
    const msg=e.data;
    if(msg.startsWith('__DONE__:')){
      es.close(); btn.disabled=false;
      const ok=msg.includes(':OK');
      status.className='status '+(ok?'ok':'error');
      status.innerHTML='<span class="dot"></span>'+(ok?'Completed':'Error');
      logtgl.style.display='inline';
      if(ok){ loadResults(key); loadAnalytics(key); }
      return;
    }
    const d=document.createElement('div');
    const l=msg.toLowerCase();
    if(l.includes('error')) d.className='line-err';
    else if(l.startsWith('  ')) d.className='line-inf';
    else if(l.includes('downloaded')||l.includes('generated')||l.includes('completed')||l.includes('successfully')) d.className='line-ok';
    d.textContent=msg; term.appendChild(d); term.scrollTop=term.scrollHeight;
  };
  es.onerror=function(){
    es.close(); btn.disabled=false;
    status.className='status error';
    status.innerHTML='<span class="dot"></span>No connection';
  };
}

function loadResults(key){
  fetch('/results/'+key).then(r=>r.json()).then(data=>{
    const panel=document.getElementById('panel-'+key);
    const tbody=document.getElementById('tbody-'+key);
    const badge=document.getElementById('badge-'+key);
    const section=document.getElementById('results-section');

    if(data && data.error){
      badge.textContent='Read error';
      badge.className='panel-badge';
      panel.classList.add('visible');
      section.classList.add('visible');
      return;
    }
    if(!Array.isArray(data)||!data.length){
      badge.textContent='No open items';
      badge.className='panel-badge all-clear';
      panel.classList.add('visible');
      section.classList.add('visible');
      return;
    }
    tbody.innerHTML='';
    data.forEach(grp=>{
      const hdr=document.createElement('tr');
      hdr.className='acc-hdr';
      hdr.innerHTML=
        '<td class="acc-num">'+grp.account+'</td>'+
        '<td class="acc-cc">'+grp.cc+'</td>'+
        '<td colspan="3"></td>'+
        '<td class="acc-bal"><span style="background:#EDFBEE;color:#188918;padding:2px 10px;border-radius:100px;font-size:.75rem;font-weight:700;">✓ CLEARABLE</span></td>';
      tbody.appendChild(hdr);
      grp.documents.forEach(doc=>{
        const tr=document.createElement('tr');
        tr.className='doc-row';
        tr.innerHTML=
          '<td></td>'+
          '<td></td>'+
          '<td>'+doc.doc+'</td>'+
          '<td>'+doc.ym+'</td>'+
          '<td class="amt '+(doc.amount<0?'neg':'pos')+'">'+fmt(doc.amount)+'</td>'+
          '<td class="txt">'+(doc.text||'')+'</td>';
        tbody.appendChild(tr);
      });
    });
    badge.textContent=data.length+(data.length===1?' account':' accounts');
    badge.className='panel-badge';
    panel.classList.add('visible');
    section.classList.add('visible');
  }).catch(()=>{
    const badge=document.getElementById('badge-'+key);
    badge.textContent='File not found';
    badge.className='panel-badge';
    document.getElementById('panel-'+key).classList.add('visible');
    document.getElementById('results-section').classList.add('visible');
  });
}
// ── Analytics ────────────────────────────────────────────────
const analyticsCache={};
let activeKey=null;
let cStatus=null,cDocs=null,cPeriod=null;

// Palette — reference instance (palette.md) + status tokens
const P={
  clear:'#0ca30c', open:'#d03b3b', blue:'#2a78d6',
  grid:'#e1e0d9', muted:'#898781', sub:'#52514e', surface:'#ffffff'
};
const BASE_OPT={
  responsive:true, maintainAspectRatio:false,
  animation:{duration:300},
  plugins:{tooltip:{cornerRadius:4,padding:8,
    callbacks:{label:ctx=>' '+ctx.dataset.label+': '+ctx.formattedValue}}}
};

function fetchAnalytics(key){
  return fetch('/summary/'+key).then(r=>r.json()).then(d=>{
    if(!d||d.error||Array.isArray(d)) return false;
    analyticsCache[key]=d; return true;
  }).catch(()=>false);
}

function showAnalytics(key){
  activeKey=key;
  document.getElementById('analytics-section').style.display='block';
  document.querySelectorAll('.flow-tab').forEach(b=>b.classList.remove('active'));
  const t=document.getElementById('atab-'+key); if(t) t.classList.add('active');
  renderAnalytics(key);
}

function loadAnalytics(key){
  if(analyticsCache[key]){showAnalytics(key);return;}
  fetchAnalytics(key).then(ok=>{if(ok) showAnalytics(key);});
}

function switchAnalytics(key,btn){
  activeKey=key;
  document.querySelectorAll('.flow-tab').forEach(b=>b.classList.remove('active'));
  btn.classList.add('active');
  if(analyticsCache[key]) renderAnalytics(key);
  else fetchAnalytics(key).then(ok=>{if(ok) renderAnalytics(key);});
  if(_rdata[key]){
    ['us','canada','concur'].forEach(k=>{
      document.getElementById('panel-'+k).classList.toggle('visible', k===key);
    });
    document.getElementById('results-section').classList.add('visible');
  } else {
    loadResults(key);
  }
}

function renderAnalytics(key){
  const d=analyticsCache[key]; if(!d) return;
  // KPIs
  document.getElementById('kpi-total').textContent=d.total_combos;
  document.getElementById('kpi-clear').textContent=d.clear_combos+' ('+d.clear_pct+'%)';
  document.getElementById('kpi-clear-sub').textContent=
    d.open_combos>0?'of '+d.total_combos+' combinations':'All accounts at zero';
  document.getElementById('kpi-open').textContent=d.open_combos;
  document.getElementById('kpi-docs').textContent=d.total_docs_clearable;
  // Charts
  buildStatusChart(d.by_account);
  buildDocsChart(d.docs_by_account);
  if(d.docs_by_period&&d.docs_by_period.length>0){
    document.getElementById('period-card').style.display='block';
    buildPeriodChart(d.docs_by_period);
  } else {
    document.getElementById('period-card').style.display='none';
  }
}

function dynHeight(rows,rowPx,min,max){
  return Math.min(max,Math.max(min,rows*rowPx+60));
}

function hAxis(){ return {
  grid:{color:P.grid,lineWidth:1},border:{display:false},
  ticks:{color:P.muted,font:{size:11},precision:0}
};}
function vAxis(){ return {
  grid:{display:false},border:{display:false},
  ticks:{color:P.sub,font:{size:11}}
};}

function buildStatusChart(byAcc){
  const wrap=document.getElementById('wrap-status');
  wrap.style.height=dynHeight(byAcc.length,28,160,420)+'px';
  if(cStatus){cStatus.destroy();cStatus=null;}
  cStatus=new Chart(document.getElementById('chart-status'),{
    type:'bar',
    data:{
      labels:byAcc.map(d=>d.account),
      datasets:[
        {label:'Clearable',data:byAcc.map(d=>d.clear),
         backgroundColor:P.clear,maxBarThickness:20,
         borderRadius:{topRight:4,bottomRight:4},borderWidth:2,borderColor:P.surface},
        {label:'Open Items',data:byAcc.map(d=>d.open),
         backgroundColor:P.open,maxBarThickness:20,
         borderRadius:{topRight:4,bottomRight:4},borderWidth:2,borderColor:P.surface}
      ]
    },
    options:{...BASE_OPT,
      indexAxis:'y',
      plugins:{...BASE_OPT.plugins,
        legend:{position:'top',align:'end',
          labels:{boxWidth:12,boxHeight:12,font:{size:11},color:P.sub,padding:14}}},
      scales:{x:{...hAxis(),stacked:true},y:{...vAxis(),stacked:true}}
    }
  });
}

function buildDocsChart(docsAcc){
  const sorted=[...docsAcc].sort((a,b)=>a.count-b.count);
  const wrap=document.getElementById('wrap-docs');
  wrap.style.height=dynHeight(sorted.length,28,160,420)+'px';
  if(cDocs){cDocs.destroy();cDocs=null;}
  cDocs=new Chart(document.getElementById('chart-docs'),{
    type:'bar',
    data:{
      labels:sorted.map(d=>d.account),
      datasets:[{label:'Documents',data:sorted.map(d=>d.count),
        backgroundColor:P.blue,maxBarThickness:20,
        borderRadius:{topRight:4,bottomRight:4},borderWidth:0}]
    },
    options:{...BASE_OPT,
      indexAxis:'y',
      plugins:{...BASE_OPT.plugins,legend:{display:false}},
      scales:{x:hAxis(),y:vAxis()}
    }
  });
}

function buildPeriodChart(byPeriod){
  if(cPeriod){cPeriod.destroy();cPeriod=null;}
  cPeriod=new Chart(document.getElementById('chart-period'),{
    type:'bar',
    data:{
      labels:byPeriod.map(d=>d.ym),
      datasets:[{label:'Documents',data:byPeriod.map(d=>d.count),
        backgroundColor:P.blue,maxBarThickness:24,
        borderRadius:{topLeft:4,topRight:4},borderWidth:0}]
    },
    options:{...BASE_OPT,
      plugins:{...BASE_OPT.plugins,legend:{display:false}},
      scales:{
        x:{grid:{display:false},border:{display:false},ticks:{color:P.muted,font:{size:11}}},
        y:{grid:{color:P.grid,lineWidth:1},border:{display:false},
           ticks:{color:P.muted,font:{size:11},precision:0}}
      }
    }
  });
}


// ── Clearing ──────────────────────────────────────────────────────────
function startClearAll(key){
  const data=_rdata[key];
  if(!data||!data.length)return;
  _pend={key,accounts:data.map(g=>({account:g.account,cc:g.cc,documents:dedupDocs(g.documents)}))};
  _openModal();
}

function startClearSingleBtn(btn){
  startClearSingle(btn.dataset.key,btn.dataset.acc,btn.dataset.cc);
}

function startClearSingle(key,account,cc){
  const data=_rdata[key]; if(!data)return;
  const grp=data.find(g=>g.account===account&&g.cc===cc); if(!grp)return;
  _pend={key,accounts:[{account,cc,documents:dedupDocs(grp.documents)}]};
  _openModal();
}

function dedupDocs(docs){
  const seen=new Set();
  const out=[];
  for(const d of docs){if(!seen.has(d.doc)){seen.add(d.doc);out.push(d.doc);}}
  return out;
}

function _openModal(){
  const labels={us:'&#127482;&#127480; United States (ISP)',canada:'&#127464;&#127462; Canada (ISP)',concur:'&#128203; Concur (I4P)'};
  const m=document.getElementById('clear-modal');
  const n=(_pend&&_pend.accounts)?_pend.accounts.length:0;
  m.querySelector('.modal-hdr span').textContent=
    'Clear -- '+(labels[_pend&&_pend.key]||(_pend&&_pend.key)||'?')+
    ' - '+n+' account'+(n!==1?'s':'');
  m.style.display='flex';
  document.getElementById('clear-date-inp').value=new Date().toISOString().slice(0,10);
  const log=document.getElementById('clear-modal-log');
  log.innerHTML=''; log.classList.remove('on');
  document.getElementById('clear-ok-btn').disabled=false;
  document.getElementById('clear-cancel-btn').disabled=false;
  document.getElementById('clear-resume-btn').style.display='none';
}

function _closeModal(){
  document.getElementById('clear-modal').style.display='none';
  _pend=null;
}

function confirmClear(){
  if(!_pend)return;
  const raw=document.getElementById('clear-date-inp').value;
  if(!raw){alert('Select a clearing date.');return;}
  const [yy,mm,dd]=raw.split('-');
  const sapDate=mm+'/'+dd+'/'+yy;
  document.getElementById('clear-ok-btn').disabled=true;
  document.getElementById('clear-cancel-btn').disabled=true;
  const log=document.getElementById('clear-modal-log');
  log.innerHTML=''; log.classList.add('on');
  const key=_pend.key;

  fetch('/clear/'+key,{
    method:'POST',
    headers:{'Content-Type':'application/json'},
    body:JSON.stringify({clearing_date:sapDate,accounts:_pend.accounts})
  }).then(r=>r.json()).then(({task_id})=>{
    if(!task_id)throw new Error('No task_id returned');
    _currentTaskId=task_id;
    const es=new EventSource('/clear-stream/'+task_id);
    es.onmessage=function(e){
      const msg=e.data;
      if(msg.startsWith('__PAUSE__:')){
        const sapMsg=msg.slice(10);
        document.getElementById('clear-resume-btn').style.display='';
        document.getElementById('clear-cancel-btn').disabled=false;
        const d1=document.createElement('div');
        d1.textContent='WARNING SAP rejected: '+sapMsg;
        d1.style.color='#EF5350';
        log.appendChild(d1);
        const d2=document.createElement('div');
        d2.textContent='Double-click the amount in SAP, then press Continue';
        d2.style.color='#FFB74D';
        log.appendChild(d2);
        log.scrollTop=log.scrollHeight;
        return;
      }
      if(msg.startsWith('__DONE_CLEAR__')){
        es.close();
        _currentTaskId=null;
        document.getElementById('clear-resume-btn').style.display='none';
        document.getElementById('clear-cancel-btn').disabled=false;
        fetch('/clearing-results/'+key).then(r=>r.json()).then(results=>{
          results.forEach(r=>{
            const docEl=document.getElementById('cdoc-'+r.account+'-'+r.cc);
            const btnEl=document.getElementById('cbtn-'+r.account+'-'+r.cc);
            if(docEl){
              docEl.style.display='';
              if(r.status==='ok'){docEl.textContent='Doc: '+r.clearing_doc;docEl.className='cdoc cdoc-ok';}
              else{docEl.textContent='Error';docEl.title=r.error;docEl.className='cdoc cdoc-err';}
            }
            if(btnEl)btnEl.disabled=true;
          });
          const cab=document.getElementById('clearall-'+key);
          if(cab)cab.disabled=true;
        });
        const isOk=msg.includes(':OK');
        const fin=document.createElement('div');
        fin.textContent=isOk?'Completed':'Finished with errors';
        fin.style.color=isOk?'#66BB6A':'#EF5350';
        fin.style.fontWeight='700';
        log.appendChild(fin); log.scrollTop=log.scrollHeight;
        if(isOk)setTimeout(_closeModal,2500);
        return;
      }
      document.getElementById('clear-resume-btn').style.display='none';
      const div=document.createElement('div');
      const l=msg.toLowerCase();
      if(l.includes('ok ')||l.includes('saved')||l.includes('completed'))div.style.color='#66BB6A';
      else if(l.includes('error')||l.includes('warning'))div.style.color='#EF5350';
      else if(l.includes('retry')||l.includes('->'))div.style.color='#42A5F5';
      div.textContent=msg; log.appendChild(div); log.scrollTop=log.scrollHeight;
    };
    es.onerror=function(){
      es.close();
      _currentTaskId=null;
      document.getElementById('clear-resume-btn').style.display='none';
      document.getElementById('clear-cancel-btn').disabled=false;
      document.getElementById('clear-ok-btn').disabled=false;
    };
  }).catch(err=>{
    document.getElementById('clear-ok-btn').disabled=false;
    document.getElementById('clear-cancel-btn').disabled=false;
    const div=document.createElement('div');
    div.textContent='Error: '+err; div.style.color='#EF5350';
    log.appendChild(div); log.classList.add('on');
  });
}

function resumeClear(){
  if(!_currentTaskId)return;
  document.getElementById('clear-resume-btn').style.display='none';
  document.getElementById('clear-cancel-btn').disabled=true;
  const div=document.createElement('div');
  div.textContent='Resuming clearing...';
  div.style.color='#42A5F5';
  document.getElementById('clear-modal-log').appendChild(div);
  document.getElementById('clear-modal-log').scrollTop=999999;
  fetch('/clear-resume/'+_currentTaskId,{method:'POST'});
}

// Auto-load on page load if pivot files exist
window.addEventListener('load',function(){
  ['us','canada','concur'].forEach(k=>{
    fetchAnalytics(k).then(ok=>{if(ok&&!activeKey) showAnalytics(k);});
  });
});
</script>
</body>
</html>"""


@app.route("/")
def index():
    return render_template_string(HTML)


@app.route("/run/<key>")
def run_script(key):
    if key not in SCRIPTS:
        return "Script no encontrado", 404

    script_path = os.path.join(SCRIPT_DIR, SCRIPTS[key])

    def generate():
        proc = subprocess.Popen(
            [sys.executable, script_path],
            cwd=SCRIPT_DIR,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        for line in proc.stdout:
            yield f"data: {line.rstrip()}\n\n"
        proc.wait()
        status = "OK" if proc.returncode == 0 else f"ERROR:{proc.returncode}"
        yield f"data: __DONE__:{status}\n\n"

    return Response(
        generate(),
        mimetype="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.route("/results/<key>")
def get_results(key):
    if key not in OUTPUT_FILES:
        return jsonify([]), 404
    pivot_path = os.path.join(SCRIPT_DIR, OUTPUT_FILES[key])
    if not os.path.exists(pivot_path):
        return jsonify([]), 404
    try:
        def ns(v):
            try:
                f = float(v)
                return str(int(f)) if f == int(f) else str(f)
            except:
                return str(v).strip()

        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False).name
        shutil.copy2(pivot_path, tmp)
        try:
            docs_df = pd.read_excel(tmp, sheet_name="Docs a Clearear")
        finally:
            os.unlink(tmp)

        # Detectar columna de importe (puede llamarse "Amount in LC" o "Amount in Local Currency")
        amt_col = next(
            (c for c in docs_df.columns if "amount" in c.lower()),
            None
        )
        has_ym = "Year/Month" in docs_df.columns

        result = []
        for (acc, cc), group in docs_df.groupby(["G/L Account", "Company Code"]):
            doc_list = []
            sort_cols = ["Document Number"] if "Document Number" in group.columns else []
            if sort_cols:
                group = group.sort_values(sort_cols)
            for _, row in group.iterrows():
                raw_amt = row.get(amt_col, 0) if amt_col else 0
                try:
                    amt = round(float(raw_amt), 2)
                except (TypeError, ValueError):
                    amt = 0.0
                doc_list.append({
                    "doc":    str(row.get("Document Number", "")),
                    "amount": amt,
                    "text":   str(row.get("Text", "")) if pd.notna(row.get("Text", "")) else "",
                    "ym":     str(row.get("Year/Month", "")) if has_ym and pd.notna(row.get("Year/Month", "")) else "",
                    "assign": str(row.get("Assignment", "")) if pd.notna(row.get("Assignment", "")) else "",
                })
            result.append({
                "account":   ns(acc),
                "cc":        ns(cc),
                "balance":   0.0,
                "documents": doc_list,
            })
        result.sort(key=lambda x: (x["account"], x["cc"]))
        return jsonify(result)
    except Exception as e:
        import traceback
        return jsonify({"error": str(e), "detail": traceback.format_exc()}), 500


@app.route("/summary/<key>")
def get_summary(key):
    if key not in OUTPUT_FILES:
        return jsonify({"error": "not found"}), 404
    pivot_path = os.path.join(SCRIPT_DIR, OUTPUT_FILES[key])
    if not os.path.exists(pivot_path):
        return jsonify({"error": "no file"}), 404
    try:
        def ns(v):
            try:
                f = float(v)
                return str(int(f)) if f == int(f) else str(f)
            except:
                return str(v).strip()

        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False).name
        shutil.copy2(pivot_path, tmp)
        try:
            pivot_df = pd.read_excel(tmp, sheet_name="Pivot Clearing")
            docs_df  = pd.read_excel(tmp, sheet_name="Docs a Clearear")
        finally:
            os.unlink(tmp)

        # KPI counts
        total     = len(pivot_df)
        clear_ct  = int((pivot_df["Balance"] == 0).sum())
        open_ct   = total - clear_ct
        clear_pct = round(clear_ct / total * 100, 1) if total > 0 else 0.0

        # By-account aggregation for status chart
        pivot_df["_acc"] = pivot_df["G/L Account"].apply(ns)
        by_account = []
        for acc, grp in pivot_df.groupby("_acc"):
            by_account.append({
                "account": acc,
                "clear":   int((grp["Balance"] == 0).sum()),
                "open":    int((grp["Balance"] != 0).sum()),
            })
        by_account.sort(key=lambda x: x["account"])

        # Docs by account (clearable)
        amt_col = next((c for c in docs_df.columns if "amount" in c.lower()), None)
        docs_by_acc = []
        if "G/L Account" in docs_df.columns:
            for acc, grp in docs_df.groupby("G/L Account"):
                amt = 0.0
                if amt_col:
                    try: amt = round(float(grp[amt_col].abs().sum()), 2)
                    except: pass
                docs_by_acc.append({"account": ns(acc), "count": len(grp), "amount": amt})
            docs_by_acc.sort(key=lambda x: x["count"], reverse=True)
            docs_by_acc = docs_by_acc[:15]

        # Docs by period
        docs_by_period = []
        if "Year/Month" in docs_df.columns:
            for ym, grp in sorted(
                docs_df.dropna(subset=["Year/Month"]).groupby("Year/Month"),
                key=lambda x: str(x[0])
            ):
                amt = 0.0
                if amt_col:
                    try: amt = round(float(grp[amt_col].abs().sum()), 2)
                    except: pass
                docs_by_period.append({"ym": str(ym), "count": len(grp), "amount": amt})

        return jsonify({
            "total_combos":         total,
            "clear_combos":         clear_ct,
            "open_combos":          open_ct,
            "clear_pct":            clear_pct,
            "total_docs_clearable": len(docs_df),
            "by_account":           by_account,
            "docs_by_account":      docs_by_acc,
            "docs_by_period":       docs_by_period,
        })
    except Exception as e:
        import traceback
        return jsonify({"error": str(e), "detail": traceback.format_exc()}), 500


@app.route("/clear/<key>", methods=["POST"])
def clear_start(key):
    if key not in SCRIPTS:
        return jsonify({"error": "not found"}), 404
    data = request.json or {}
    task_id = f"{key}_{int(time.time() * 1000)}"
    _clear_tasks[task_id] = {"key": key, "q": queue.Queue()}
    t = threading.Thread(target=_run_clearing_task, args=(task_id, key, data), daemon=True)
    t.start()
    return jsonify({"task_id": task_id})


@app.route("/clear-stream/<task_id>")
def clear_stream(task_id):
    if task_id not in _clear_tasks:
        return "Task not found", 404
    q_ref = _clear_tasks[task_id]["q"]

    def generate():
        while True:
            try:
                msg = q_ref.get(timeout=300)
                yield f"data: {msg}\n\n"
                if msg.startswith("__DONE_CLEAR__"):
                    break
            except queue.Empty:
                yield "data: ERROR: Timeout waiting for SAP\n\n"
                yield "data: __DONE_CLEAR__:ERROR\n\n"
                break

    return Response(
        generate(),
        mimetype="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.route("/clearing-results/<key>")
def get_clearing_results_api(key):
    return jsonify(_clearing_results.get(key, []))


@app.route("/clear-resume/<task_id>", methods=["POST"])
def clear_resume(task_id):
    if task_id in _clear_pause_events:
        _clear_pause_events[task_id].set()
        return jsonify({"ok": True})
    return jsonify({"error": "task not paused"}), 404


@app.route("/download-open-items/<key>")
def download_open_items(key):
    if key not in INPUT_FILES:
        return jsonify({"error": "not found"}), 404
    input_path = os.path.join(SCRIPT_DIR, INPUT_FILES[key])
    pivot_path = os.path.join(SCRIPT_DIR, OUTPUT_FILES[key])
    if not os.path.exists(input_path):
        return jsonify({"error": "Input file not found. Run the script first."}), 404
    try:
        if key == "concur":
            sheet_name = "Data"
        elif key == "canada":
            with pd.ExcelFile(input_path) as xl:
                sheet_name = "Data" if "Data" in xl.sheet_names else xl.sheet_names[0]
        else:
            sheet_name = None

        df = pd.read_excel(input_path, sheet_name=sheet_name, header=5)

        if key == "concur" and "Clearing Document" in df.columns:
            df = df[df["Clearing Document"].isna()]
        if "Year/Month" in df.columns:
            df = df[df["Year/Month"] >= "2024/01"]

        clearable_docs = set()
        if os.path.exists(pivot_path):
            try:
                with pd.ExcelFile(pivot_path) as xl2:
                    docs_df = pd.read_excel(xl2, sheet_name="Docs a Clearear")
                clearable_docs = set(
                    str(int(float(d))) for d in docs_df["Document Number"].dropna()
                )
            except Exception:
                pass

        doc_col = next((c for c in df.columns if "document" in c.lower() and "number" in c.lower()), None)
        if clearable_docs and doc_col:
            def _docstr(v):
                try:
                    return str(int(float(v)))
                except Exception:
                    return str(v)
            df = df[~df[doc_col].apply(_docstr).isin(clearable_docs)]

        wb = openpyxl.load_workbook(input_path)
        ws = wb[sheet_name] if sheet_name else wb.active
        if ws.max_row >= 7:
            ws.delete_rows(7, ws.max_row - 6)
        header_cols = [c.value for c in ws[6]]
        for _, row in df.iterrows():
            row_values = []
            for col_name in header_cols:
                if col_name is None:
                    row_values.append(None)
                    continue
                val = row.get(col_name)
                try:
                    row_values.append(None if pd.isna(val) else val)
                except Exception:
                    row_values.append(val)
            ws.append(row_values)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"Open_Items_{key.upper()}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(
            buf,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    except Exception as e:
        import traceback
        return jsonify({"error": str(e), "detail": traceback.format_exc()}), 500


if __name__ == "__main__":
    import webbrowser
    print("Iniciando SAP Clearing Dashboard en http://localhost:5001")
    webbrowser.open("http://localhost:5001")
    app.run(host="0.0.0.0", port=5001, threaded=True)